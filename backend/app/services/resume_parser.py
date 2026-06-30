"""Resume parsing for PDF and DOCX files.

Extracts raw text and a best-effort structured profile (name, email, phone,
location, links, skills, titles, years of experience). The heuristics are
deliberately conservative – everything extracted can be edited by the user in
the UI before it is used.
"""

from __future__ import annotations

import io
import re

import pdfplumber
from docx import Document

from ..models import CandidateProfile, ContactLinks

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(
    r"(?:(?:\+?\d{1,3}[\s.\-]?)?(?:\(?\d{2,4}\)?[\s.\-]?)?\d{3,4}[\s.\-]?\d{3,4})"
)
LINKEDIN_RE = re.compile(r"(https?://)?(www\.)?linkedin\.com/[A-Za-z0-9_\-/%.]+", re.I)
GITHUB_RE = re.compile(r"(https?://)?(www\.)?github\.com/[A-Za-z0-9_\-/.]+", re.I)
URL_RE = re.compile(r"https?://[^\s)>\]]+", re.I)
YEARS_RE = re.compile(r"(\d{1,2})\+?\s*(?:years|yrs)\b", re.I)

# A pragmatic, extensible technology / skill dictionary. Matching is
# case-insensitive and word-boundary aware.
SKILL_VOCAB = [
    "python", "java", "javascript", "typescript", "go", "golang", "rust", "c++",
    "c#", "ruby", "php", "scala", "kotlin", "swift", "objective-c", "sql", "r",
    "react", "next.js", "vue", "angular", "svelte", "node.js", "express",
    "django", "flask", "fastapi", "spring", "rails", "laravel", ".net",
    "postgresql", "mysql", "mongodb", "redis", "elasticsearch", "cassandra",
    "dynamodb", "snowflake", "bigquery", "kafka", "rabbitmq", "spark", "hadoop",
    "airflow", "dbt", "aws", "azure", "gcp", "docker", "kubernetes", "terraform",
    "ansible", "jenkins", "github actions", "gitlab ci", "ci/cd", "linux",
    "graphql", "rest", "grpc", "microservices", "machine learning",
    "deep learning", "nlp", "computer vision", "pytorch", "tensorflow",
    "scikit-learn", "pandas", "numpy", "data analysis", "data engineering",
    "html", "css", "tailwind", "sass", "figma", "agile", "scrum", "jira",
    "git", "bash", "powershell", "selenium", "playwright", "cypress", "jest",
    "pytest", "openai", "llm", "langchain", "rag", "prometheus", "grafana",
    "datadog", "kafka", "salesforce", "tableau", "power bi", "excel",
]

TITLE_KEYWORDS = [
    "engineer", "developer", "scientist", "analyst", "manager", "designer",
    "architect", "consultant", "lead", "director", "administrator", "specialist",
    "researcher", "intern", "devops", "sre", "product", "data", "frontend",
    "backend", "full stack", "fullstack", "mobile", "qa", "security",
]


def _read_pdf(data: bytes) -> str:
    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts)


def _read_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _read_pdf(data)
    if lower.endswith(".docx"):
        return _read_docx(data)
    if lower.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")


def _guess_name(lines: list[str], email: str | None) -> str | None:
    # The name is usually one of the first non-empty lines: 2-4 words, mostly
    # alphabetic, not containing an @ or digits.
    for line in lines[:8]:
        cleaned = line.strip()
        if not cleaned or "@" in cleaned or any(ch.isdigit() for ch in cleaned):
            continue
        words = cleaned.split()
        if 1 < len(words) <= 4 and all(
            w.replace("-", "").replace(".", "").isalpha() for w in words
        ):
            if cleaned.lower() not in {"curriculum vitae", "resume"}:
                return cleaned.title() if cleaned.isupper() else cleaned
    if email:
        local = email.split("@")[0]
        parts = re.split(r"[._\-]", local)
        if len(parts) >= 2:
            return " ".join(p.capitalize() for p in parts if p)
    return None


def _normalize_url(value: str) -> str:
    value = value.rstrip(".,;)")
    if not value.lower().startswith("http"):
        value = "https://" + value
    return value


def _find_skills(text: str) -> list[str]:
    lower = text.lower()
    found: list[str] = []
    for skill in SKILL_VOCAB:
        pattern = r"(?<![A-Za-z0-9])" + re.escape(skill.lower()) + r"(?![A-Za-z0-9])"
        if re.search(pattern, lower):
            found.append(skill)
    # de-dupe while preserving order
    seen: set[str] = set()
    out: list[str] = []
    for s in found:
        if s not in seen:
            seen.add(s)
            out.append(s)
    return out


def _find_titles(lines: list[str], name: str | None = None) -> list[str]:
    titles: list[str] = []
    name_l = name.lower().strip() if name else None
    for line in lines:
        low = line.lower().strip()
        if not low or len(low) > 60:
            continue
        # skip contact lines and the candidate's own name line
        if "@" in low or "|" in low or any(ch.isdigit() for ch in low):
            continue
        if name_l and low == name_l:
            continue
        if any(kw in low for kw in TITLE_KEYWORDS):
            # avoid sentences
            if len(low.split()) <= 6:
                cleaned = line.strip(" -•\t")
                if cleaned and cleaned not in titles:
                    titles.append(cleaned)
        if len(titles) >= 6:
            break
    return titles


def _guess_location(text: str) -> str | None:
    # Look for "City, ST" or "City, Country" patterns near the top.
    head = "\n".join(text.splitlines()[:15])
    m = re.search(r"\b([A-Z][a-zA-Z]+(?:\s[A-Z][a-zA-Z]+)?),\s*([A-Z]{2}|[A-Z][a-zA-Z]+)\b", head)
    if m:
        return m.group(0)
    return None


def parse_resume(filename: str, data: bytes) -> CandidateProfile:
    text = extract_text(filename, data)
    lines = [ln for ln in text.splitlines()]
    non_empty = [ln.strip() for ln in lines if ln.strip()]

    email_match = EMAIL_RE.search(text)
    email = email_match.group(0) if email_match else None

    phone = None
    for m in PHONE_RE.finditer(text):
        candidate = m.group(0)
        digits = re.sub(r"\D", "", candidate)
        if 7 <= len(digits) <= 15:
            phone = candidate.strip()
            break

    linkedin = None
    li = LINKEDIN_RE.search(text)
    if li:
        linkedin = _normalize_url(li.group(0))

    github = None
    gh = GITHUB_RE.search(text)
    if gh:
        github = _normalize_url(gh.group(0))

    website = None
    for u in URL_RE.finditer(text):
        url = u.group(0)
        if "linkedin.com" in url.lower() or "github.com" in url.lower():
            continue
        website = _normalize_url(url)
        break

    skills = _find_skills(text)
    name = _guess_name(non_empty, email)
    titles = _find_titles(non_empty, name)
    location = _guess_location(text)

    years = None
    year_matches = [int(m.group(1)) for m in YEARS_RE.finditer(text)]
    if year_matches:
        years = float(max(year_matches))

    headline = titles[0] if titles else None
    summary = None
    # crude summary = first paragraph after a "summary"/"about" header
    for i, ln in enumerate(non_empty):
        if re.match(r"^(summary|about|profile|objective)\b", ln, re.I):
            summary = " ".join(non_empty[i + 1 : i + 4])
            break

    return CandidateProfile(
        full_name=name,
        email=email,
        phone=phone,
        location=location,
        headline=headline,
        summary=summary,
        links=ContactLinks(linkedin=linkedin, github=github, website=website),
        skills=skills,
        titles=titles,
        years_experience=years,
        resume_filename=filename,
        resume_text=text,
        raw_resume_excerpt=text[:1500],
    )
